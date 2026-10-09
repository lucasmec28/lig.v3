"""Verificações independentes do caminho da força e das hipóteses da versão."""
from dataclasses import replace
from io import BytesIO
import math
import pytest
from docx import Document
from lro.models import Connection,Profile,profiles
from lro.examples import presets
from lro.engine import evaluate
from lro.column_checks import normal_patch,column_actions,crippling_point,column_checks
from lro.report import create_report


def column_case():
    return next(c for c in presets().values() if c.support.name=='CS 600 x 281')


def test_linear_patch_hand_solution_and_conservation():
    # N=20 kN; M=2,2 kNm; h=220 mm → p+=363,636 e p−=−181,818 N/mm.
    # Zona comprimida triangular de 73,333 mm; C=6,6667 kN; T=26,6667 kN.
    p=normal_patch(20000,2200000,220)
    assert p.peak==pytest.approx(4000/11)
    assert p.low==pytest.approx(-2000/11)
    assert p.compression_length==pytest.approx(220/3)
    assert p.compression==pytest.approx(20000/3)
    assert p.tension==pytest.approx(80000/3)
    assert p.tension-p.compression==pytest.approx(20000)


@pytest.mark.parametrize('N,M',[(0,2200000),(20000,2200000),(50000,10000),(45000,0)])
def test_patch_matches_numerical_integration(N,M):
    h=220.;dy=h/20000;force=torque=positive=negative=0.
    for i in range(20000):
        y=-h/2+(i+.5)*dy;q=N/h+12*M*y/h**3
        force+=q*dy;torque+=q*y*dy
        positive+=max(q,0)*dy;negative+=max(-q,0)*dy
    p=normal_patch(N,M,h)
    assert force==pytest.approx(N,abs=1e-6)
    assert torque==pytest.approx(M,abs=.01)
    assert p.tension==pytest.approx(positive,abs=.001)
    assert p.compression==pytest.approx(negative,abs=.001)
    assert normal_patch(N,-M,h)==p


def test_column_transfer_includes_flange_thickness_and_sign_envelope():
    c=column_case();M,p,q=column_actions(c,11000,2000)
    assert M==pytest.approx(1071400) # 11000*(75+22,4); chapa centrada.
    assert M-11000*c.bolt_centroid_x==pytest.approx(246400)
    assert column_actions(c,-11000,2000)==(M,p,q)


def test_nbr_574_point_end_case_hand_calculation():
    # tw=10, tf=20, fy=250 → sqrt(E fy tf/tw)=10000 MPa; 0,33*100/1,10=30 mm².
    assert crippling_point(10,20,250)==pytest.approx(300000)


def test_nbr_573_column_point_yield_result():
    c=column_case()
    s=replace(c.support,tw=10,tf=20)
    c=replace(c,support=s,a=80)
    rows={x.id:x for x in column_checks(c,22000,20000,'Teste manual')}
    # M=22000*(80+20)=2,2e6 → T=26666,667 N; FRd=2,5*20*250*10.
    assert rows['support_web_y'].demand==pytest.approx(80000/3)
    assert rows['support_web_y'].resistance==pytest.approx(125000)
    assert rows['support_crippling'].demand==pytest.approx(20000/3)


def test_user_column_with_cjp_is_complete_within_explicit_assumptions():
    c=column_case();r=evaluate(c)
    assert (c.V,c.N,c.gap,c.a,c.plate_top)==(11000,2000,10,75,89.5)
    assert r.status=='ATENDE AO ESCOPO VERIFICADO'
    assert r.governing.ratio==pytest.approx(.45393683714510796)
    assert {'support_joint_base','support_web_y','support_crippling'}<={x.id for x in r.checks}
    assert not any(x.id=='support_joint_weld' for x in r.checks)
    assert c.fixed_assumptions['welded_profile_internal_joint'].startswith('complete_joint_penetration')


def test_old_json_fixed_assumptions_do_not_change_geometry_and_loads():
    c=column_case();d=c.to_dict();d['schema_version']=3;d['connection']['restrained']=False
    d['connection']['support_joint_weld']=4
    restored=Connection.from_dict(d)
    assert restored.restrained is True
    assert (restored.a,restored.gap,restored.V,restored.N)==(75,10,11000,2000)
    assert evaluate(restored).checks==evaluate(c).checks
    assert restored.to_dict()['fixed_assumptions']==c.fixed_assumptions


def test_column_report_distinguishes_cjp_from_connection_fillets():
    c=column_case();d=Document(BytesIO(create_report(c,evaluate(c))))
    text='\n'.join([p.text for p in d.paragraphs]+[cell.text for t in d.tables for row in t.rows for cell in row.cells])
    assert 'penetração total' in text and 'Não há solda direta entre viga e pilar' in text
    assert 'eficaz por hipótese fixa' in text
    assert 'deslocamento lateral relativo' in text
    assert 'Não informada' not in text
    assert 'Sem cortante horizontal' in text or 'Não há cortante horizontal' in text


def test_removing_weak_axis_demand_does_not_approve_girder_web_interaction():
    c=next(iter(presets().values()));r=evaluate(c)
    assert c.restrained is True
    assert r.status=='ATENDE ÀS VERIFICAÇÕES REALIZADAS'
    assert any(x.severity=='excluded' and 'fora do plano' in x.text for x in r.issues)
    assert all('My=0' in x.variables for x in r.checks if x.id.startswith('interaction_'))
