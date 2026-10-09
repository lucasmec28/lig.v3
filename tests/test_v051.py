"""A exclusão autorizada limita a conclusão; não transforma falhas em aprovação."""
from dataclasses import replace
from io import BytesIO
from docx import Document
import pytest
from lro.examples import presets
from lro.models import SUPPORT_WEB_EXCLUSION
from lro.engine import evaluate
from lro.report import create_report


def test_combined_support_is_excluded_from_rows_and_has_one_report_note():
    c=next(iter(presets().values()));r=evaluate(c)
    excluded=[x for x in r.issues if x.severity=='excluded']
    assert [x.text for x in excluded]==[SUPPORT_WEB_EXCLUSION]
    assert not any(x.severity=='pending' for x in r.issues)
    assert not {'support_web_n','support_web_punch_n'} & {x.id for x in r.checks+r.actual}
    assert len(r.checks)==23
    assert r.governing.ratio==pytest.approx(.668666683954172)
    assert c.to_dict()['fixed_assumptions']['supporting_girder_web_combined_out_of_plane_interaction']=='not_verified'
    d=Document(BytesIO(create_report(c,r)))
    text='\n'.join([p.text for p in d.paragraphs]+[cell.text for t in d.tables for row in t.rows for cell in row.cells])
    assert text.count(SUPPORT_WEB_EXCLUSION)==1
    assert 'ATENDE ÀS VERIFICAÇÕES REALIZADAS' in text
    assert 'Pendências para conclusão' not in text
    assert 'Distância longitudinal livre' not in text


def test_exclusion_does_not_hide_failure_geometry_or_other_pending_gates():
    c=next(iter(presets().values()))
    assert evaluate(replace(c,V=300000)).status=='NÃO ATENDE'
    assert evaluate(replace(c,pitch=1)).status=='GEOMETRIA INVÁLIDA'
    assert evaluate(replace(c,norm_minimum=False)).status=='VERIFICAÇÃO INCOMPLETA'


def test_exclusion_scope_and_pure_centered_tension_model_remain_distinct():
    c=next(iter(presets().values()))
    assert evaluate(replace(c,support_edge_distance=10000)).status=='ATENDE ÀS VERIFICAÇÕES REALIZADAS'
    pure=replace(c,V=0,support_edge_distance=1000)
    r=evaluate(pure)
    assert r.status=='ATENDE AO ESCOPO VERIFICADO'
    assert {'support_web_n','support_web_punch_n'} <= {x.id for x in r.checks}
    assert not any(x.severity=='excluded' for x in r.issues)
    assert any(x.severity=='pending' and x.origin=='data' for x in evaluate(replace(pure,support_edge_distance=0)).issues)
    ecc=replace(pure,plate_top=pure.plate_top-10)
    assert ecc.support_web_combined_excluded
    assert any(x.severity=='excluded' for x in evaluate(ecc).issues)
    assert not replace(c,N=0).support_web_combined_excluded
    column=next(x for x in presets().values() if x.support.name=='CS 600 x 281')
    assert not column.support_web_combined_excluded
    assert evaluate(column).status=='ATENDE AO ESCOPO VERIFICADO'
