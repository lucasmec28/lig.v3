"""Memória editável: equações OMML, tabela única de verificações e desenho comum."""
from io import BytesIO
from copy import deepcopy
from datetime import datetime
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from .column_checks import WELDED_FAMILIES,column_actions
from .models import KGF,VERSION,STEELS,BOLTS
from .drawing import image_bytes
from .detailing import polygon_area,plate_outline

LINK='https://www.linkedin.com/in/lucas-oliveira-722723149/?isSelfProfile=true'
BRAND='Desenvolvido por LRO Soluções de engenharia LTDA.'


def number(x,places=2):
    return f'{x:,.{places}f}'.replace(',','§').replace('.',',').replace('§','.')

def decimal_text(text):
    """Converte decimais nas substituições; preserva itens como 6.2.5."""
    return re.sub(r'(?<![\w.])(\d+)\.(\d+)(?![\w.])',r'\1,\2',text)


def display(v,unit):
    factor,label={'N':(KGF,'kgf'),'Nmm':(KGF*1000,'kgf·m'),'N/mm':(KGF,'kgf/mm')}.get(unit,(1,unit))
    return number(v/factor,3 if unit=='—' else 2),label


def hyperlink(p,label,url):
    rel=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel)
    r=OxmlElement('w:r');rp=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'225D86');rp.append(color);r.append(rp)
    t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)


def table(doc,headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for col,width in zip(t.columns,widths):col.width=Inches(width)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','bottom','left','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    props.append(borders)
    for i,label in enumerate(headers):t.rows[0].cells[i].text=label
    rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
    for row in rows:
        for cell,value in zip(t.add_row().cells,row):cell.text=str(value)
    for ri,row in enumerate(t.rows):
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        for j,cell in enumerate(row.cells):
            cell.width=Inches(widths[j]);cell.vertical_alignment=1
            pr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE8F0' if ri==0 else ('F5F7F9' if ri%2==0 else 'FFFFFF'));pr.append(sh)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'65');el.set(qn('w:type'),'dxa');margins.append(el)
            pr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
                if j>0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.size=Pt(9);r.bold=(ri==0)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t


# OMML nativo: frações, índices, expoentes e radicais estruturados.
def mr(text):
    r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=str(text);r.append(t);return r
def group(*items):
    out=[]
    for item in items:
        if isinstance(item,list):out.extend(item)
        else:out.append(mr(item) if isinstance(item,str) else item)
    return out
def slot(tag,items):
    e=OxmlElement('m:'+tag)
    for item in group(items):e.append(deepcopy(item))
    return e
def sub(base,idx):
    e=OxmlElement('m:sSub');e.append(slot('e',group(base)));e.append(slot('sub',group(idx)));return e
def power(base,exp='2'):
    e=OxmlElement('m:sSup');e.append(slot('e',group(base)));e.append(slot('sup',group(exp)));return e
def frac(a,b):
    e=OxmlElement('m:f');e.append(slot('num',group(a)));e.append(slot('den',group(b)));return e
def radical(items):
    e=OxmlElement('m:rad');pr=OxmlElement('m:radPr');de=OxmlElement('m:degHide');de.set(qn('m:val'),'1');pr.append(de);e.append(pr);e.append(slot('deg',[]));e.append(slot('e',group(items)));return e


def native_equation(doc,check,c):
    fy=sub('f','y');fu=sub('f','u');g1=sub('γ','a1');g2=sub('γ','a2');Ag=sub('A','g');An=sub('A','n')
    cid=check.id
    if cid=='bolts':
        alpha='0,45' if c.threads or c.bolt=='ASTM A307' else '0,56'
        expr=group(sub('F','Rd'),' = ',frac(group(sub('k','pega'),alpha,'π',power(sub('d','b')),sub('f','ub')),group('4',g2)))
    elif cid.startswith('bearing'):
        expr=group(sub('F','Rd'),' = ',frac(group('min(1,2',sub('l','c'),'t',fu,'; 2,4',sub('d','b'),'t',fu,')'),g2))
    elif cid in ['plate_vy','plate_vu','plate_ny','plate_nu']:
        rupture=cid.endswith('u');shear=cid.startswith('plate_v')
        expr=group(sub('V' if shear else 'N','Rd'),' = ',frac(group('0,60' if shear else '',fu if rupture else fy,An if rupture else Ag),g2 if rupture else g1))
    elif cid.startswith('interaction'):
        expr=group('η = ',power(group('[n/2 + m]')),' + ',power('v'),' ≤ 1   (n < 0,2)')
        expr2=group('η = ',power(group('[n + 8m/9]')),' + ',power('v'),' ≤ 1   (n ≥ 0,2)')
        _math(doc,expr);_math(doc,expr2);return
    elif cid=='weld':
        expr=group(sub('q','max'),' = ',radical(group(power(group('(',frac('N','2h'),' + ',frac('3M',power('h')),')')),' + ',power(group('(',frac('V','2h'),')')))))
        _math(doc,expr)
        expr=group(sub('q','Rd'),' = ',frac(group('0,60',sub('f','w'),'w'),group(radical('2'),sub('γ','w2'))))
    elif cid=='plate_mu':expr=group(sub('M','Rd'),' = ',frac(group(fu,sub('Z','n')),g2))
    elif cid=='plate_ltb':expr=group(sub('M','Rd'),' = ',frac(sub('M','n'),g1),' ; ',sub('M','n'),' = min(',sub('M','p'),'; ',sub('C','b'),'[1,52 − 0,274λ',frac(fy,'E'),']',fy,'W)')
    elif cid=='beam_vm':expr=group('η = ',frac('M',group(sub('M','BC,Rd'),' + ',sub('V','AB,Rd'),'ℓ')),' + ',frac('N',sub('N','web,Rd')),' ≤ 1')
    elif cid=='support_punch':
        expr=group(sub('t','p'),' ≤ max[',frac(group(sub('t','s'),sub('f','u,s')),group(sub('f','y,p'),g2)),'; ',frac(group(power(sub('t','s')),sub('f','u,s'),power(sub('h','p'))),group('6Ve',g2)),']')
        if c.N>0:expr=group(sub('t','p'),' ≤ ',frac(group(sub('t','s'),sub('f','u,s')),group(sub('f','y,p'),g2)))
    elif cid=='support_shear':expr=group(sub('R','Rd'),' = min(',frac(group('1,2',fy,'h',sub('t','s')),g1),'; ',frac(group('1,2',fu,'h',sub('t','s')),g2),')')
    elif cid=='support_web_y':expr=group('max(T; C) ≤ ',frac(group('1,10(2,5k + ',sub('l','n'),')',fy,sub('t','w')),g1))
    elif cid=='support_crippling':expr=group('C ≤ ',frac(group('0,33',power(sub('t','w')),radical(group('E',fy,sub('t','f'),'/',sub('t','w')))),g1))
    elif cid=='support_joint_base':expr=group(sub('r','max'),' ≤ min(',frac(group('0,6',fy,sub('t','w')),g1),'; ',frac(group('0,6',fu,sub('t','w')),g2),')')
    elif cid=='support_flange':expr=group('T ≤ ',frac(group('0,5 × 6,25',power(sub('t','f')),fy),g1))
    elif cid.startswith('block'):
        expr=group(sub('R','Rd'),' = ',frac(group('min(0,60',fu,sub('A','nv'),' + ',fu,sub('A','nt'),'; 0,60',fy,sub('A','gv'),' + ',fu,sub('A','nt'),')'),g2))
        if cid=='block_plate':
            _math(doc,expr)
            expr=group('η = ',power(frac('V',sub('R','v'))),' + ',power(frac('N',sub('R','n'))),' ≤ 1')
    elif cid=='base_plate':expr=group(sub('q','Rd'),' = min(',frac(group('0,60',fy,sub('t','p')),g1),'; ',frac(group('0,60',fu,sub('t','p')),g2),')')
    elif cid=='beam_v':expr=group(sub('V','Rd'),' = ',frac(group('0,60',fy,sub('h','web'),sub('t','w')),g1))
    elif cid=='beam_n':expr=group(sub('N','Rd'),' = min(',frac(group(fy,Ag),g1),'; ',frac(group(fu,An),g2),')')
    elif cid=='ductility':expr=group(sub('t','p'),' ≤ ',frac(group('6',sub('R','n,par'),'C′'),group(fy,power(sub('h','p')))))
    elif cid=='weld_development':expr=group('w ≥ max(',frac(group('5',sub('t','p')),'8'),'; ',frac(group(radical('3'),sub('t','p'),fy),group('2',sub('f','w'))),')')
    elif cid=='cope_interaction':expr=group('η = ',frac('N',sub('N','Rd')),' + ',frac(sub('M','c'),sub('M','Rd')),' + ',frac('V',sub('V','Rd')),' ≤ 1')
    elif cid=='cope_rupture':expr=group('η = ',frac(sub('M','c'),group(fu,sub('Z','n'),'/',g2)),' + ',frac('N',group(fu,An,'/',g2)),' ≤ 1')
    elif cid=='cope_block':expr=group('η = ',power(frac('V',sub('R','v'))),' + ',power(frac('N',sub('R','n'))),' ≤ 1')
    elif cid=='beam_buckling':expr=group(sub('V','Rd'),' = ',frac(group('0,6',fy,'h',sub('t','w'),sub('C','v')),g1))
    elif cid=='support_web_n':
        expr=group(sub('N','Rd'),' = ',frac(group(fy,power(sub('t','w')),group('[4',radical('2wuv(u+v)'),' + L(u+v)]')),group('4',g1,'uv')))
    elif cid=='support_web_punch_n':
        expr=group(sub('N','Rd'),' = ',frac(group('0,6',fu,'·2(',sub('h','p'),' + ',sub('t','p'),')',sub('t','w')),g2))
    elif cid.startswith('full_weld_'):
        expr=group('τ = ',radical(group(power(group('(',frac('N',sub('A','w')),' − ',frac('My',sub('J','w')),')')),' + ',power(group('(',frac('V',sub('A','w')),' + ',frac('Mx',sub('J','w')),')')))),' ≤ ',frac(group('0,6',sub('f','w')),sub('γ','w2')))
    elif cid.startswith('full_base_'):
        expr=group(sub('q','Sd'),' ≤ min(',frac(group('0,6',sub('f','y,p'),sub('t','p')),g1),'; ',frac(group('0,6',sub('f','u,p'),sub('t','p')),g2),'; ',frac(group('0,6',sub('f','y,s'),sub('t','s')),g1),'; ',frac(group('0,6',sub('f','u,s'),sub('t','s')),g2),')')
    elif cid in ('full_compactness','opposite_compactness'):
        expr=group(frac(sub('b','r' if cid=='full_compactness' else 'e'),sub('t','p' if cid=='full_compactness' else 'e')),' ≤ ',frac('200',radical(fy)))
    else:
        # Novos estados limite nunca recebem uma equação de tração genérica.
        doc.add_paragraph(check.equation)
        return
    _math(doc,expr)


def _math(doc,expr):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(4)
    m=OxmlElement('m:oMath')
    for item in expr:m.append(deepcopy(item))
    p._p.append(m)


def create_report(c,r,detailed=False):
    if c.full_depth or c.opposite_stiffener:
        raise ValueError('Chapa entre mesas e enrijecedor oposto não são avaliados nesta versão.')
    doc=Document();s=doc.sections[0]
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=s.bottom_margin=Inches(.6);s.left_margin=s.right_margin=Inches(.7)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
    for style in doc.styles:
        for border in style.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    normal=doc.styles['Normal'];normal.font.size=Pt(9.5);normal.paragraph_format.space_after=Pt(5)
    doc.styles['Title'].font.size=Pt(20)
    doc.styles['Heading 1'].font.size=Pt(12);doc.styles['Heading 2'].font.size=Pt(10)
    doc.core_properties.author='LRO Soluções de engenharia LTDA.'
    doc.core_properties.title='Memória de cálculo de ligação com chapa simples'
    doc.add_paragraph('Verificações parciais da chapa entre mesas' if c.full_depth else ('Ligação com chapa simples' if r.checks else 'Pré-detalhamento de ligação'),style='Title')
    doc.add_paragraph(c.project)
    doc.add_paragraph(f'{c.beam.name} → {c.support.name} | '+('Alma da viga de apoio' if c.kind=='beam_web' else 'Mesa do pilar alinhada à alma'))
    p=doc.add_paragraph();p.add_run(r.status).bold=True
    if r.governing:
        p.add_run(f' · índice resistente máximo {number(r.governing.ratio,3)} · {r.governing.name}')
    for issue in r.issues:
        if issue.severity=='excluded':doc.add_paragraph('Nota: '+issue.text)
    pending=[i for i in r.issues if i.severity in ('pending','error')]
    if any(not x.passed for x in r.checks):
        doc.add_paragraph('A ligação não atende às verificações ou requisitos do método indicados no quadro resumo. Revisar o detalhe e conferir as pendências de escopo.')
    elif pending:
        doc.add_paragraph('Conclusão condicionada às pendências de escopo registradas abaixo. Os estados limite ainda não verificados não estão aprovados.')
    else:doc.add_paragraph('As verificações locais incluídas nesta versão atendem para a geometria e as hipóteses registradas. A análise global dos membros e os requisitos globais de integridade estrutural permanecem no projeto da estrutura.')
    if not r.checks:doc.add_paragraph('Somente geometria e materiais: não há resistência calculada nem aprovação estrutural desta configuração.')
    elif c.full_depth:
        doc.add_paragraph('SEM APROVAÇÃO DA LIGAÇÃO. Calculados somente o grupo elástico de soldas, o metal-base local e condições de compacidade. Esses índices não verificam os parafusos, a estabilidade acoplada, a capacidade de rotação ou as forças e soldas do enrijecedor oposto.')
    doc.add_picture(BytesIO(image_bytes(c)),width=Inches(7.05))
    doc.add_paragraph('Dimensões e materiais adotados',style='Heading 1')
    geom=[('Chapa',f'{number(c.width)} × {number(c.hp)} × {number(c.tp,4)} mm',c.plate_steel),('Parafusos',f'{c.n} × Ø {number(c.db,3)} mm',c.bolt),('Furos',f'Ø {number(c.dh,4)} mm', 'Padrão; '+('broca' if c.drilled else 'desconto líquido +2 mm')),('Soldas',f'2 filetes de {number(c.weld)} mm; L = {number(c.hp)} mm',f'fw = {number(c.fw,0)} MPa'),('Posições',f'z = {c.plate_top:g}; a = {c.a:g}; g = {c.gap:g}; p = {c.pitch:g}; eᵥ = eₕ = {c.edge_v:g}' if c.edge_v==c.edge_h else f'z = {c.plate_top:g}; a = {c.a:g}; g = {c.gap:g}; p = {c.pitch:g}; eᵥ = {c.edge_v:g}; eₕ = {c.edge_h:g}', 'mm'),('Viga apoiada',f'd/bf/tw/tf = {c.beam.d:g}/{c.beam.bf:g}/{c.beam.tw:g}/{c.beam.tf:g}',c.beam_steel),('Apoio',f'd/bf/tw/tf = {c.support.d:g}/{c.support.bf:g}/{c.support.tw:g}/{c.support.tf:g}',c.support_steel)]
    geom[1]=('Parafusos',f'{c.n*c.bolt_columns} × Ø {number(c.db,3)} mm; {c.n} linhas × {c.bolt_columns} coluna(s)',c.bolt)
    if c.full_depth:
        geom[0]=('Chapa recortada',f'Aba: {number(c.width)} × {number(c.hp)} mm; t = {number(c.tp,4)} mm',c.plate_steel)
        geom[3]=('Soldas',f'Alma: w = {number(c.weld)} mm; mesas: w = {number(c.flange_weld)} mm',f'fw = {number(c.fw,0)} MPa; grupo elástico isolado')
        geom.append(('Entre mesas',f'H = {number(c.root_height)}; bᵣ = {number(c.root_width)}; c = {number(c.corner_clip)} mm','Alívios a 45°; chapa soldada à alma e mesas'))
        if c.opposite_stiffener:geom.append(('Enrijecedor oposto',f'H = {number(c.root_height)}; bₑ = {number(c.stiffener_width)}; t = {number(c.stiffener_t,4)}; w = {number(c.stiffener_weld)} mm',c.stiffener_steel))
    if c.bolt_columns>1:geom.append(('Passo horizontal',f's = {number(c.gauge)} mm','Grupo bidimensional; furos alinhados'))
    if c.kind=='beam_web' and c.N>0 and not c.full_depth and not c.support_web_combined_excluded:
        geom.append(('Alma do apoio',f'Distância longitudinal livre = {number(c.support_edge_distance)} mm' if c.support_edge_distance else 'Distância longitudinal livre não informada','Menor valor nos dois sentidos a partir do eixo da chapa'))
    if c.kind=='column_flange' and c.support.family in ('CS','CVS','VS','Soldado'):
        geom.append(('Fabricação do perfil','Juntas mesa–alma de penetração total','Metal de adição compatível; hipótese de projeto'))
    table(doc,['Componente','Dimensões','Material ou especificação'],[(a,decimal_text(b),cc) for a,b,cc in geom],[1.1,3.25,2.75])
    mats=[]
    material_list=[('Chapa',c.plate_steel),('Viga',c.beam_steel),('Apoio',c.support_steel)]
    if c.opposite_stiffener:material_list.append(('Enrijecedor',c.stiffener_steel))
    for label,name in material_list:
        st=STEELS[name];mats.append(f'{label}: fy = {st.fy:g} MPa; fu = {st.fu:g} MPa')
    doc.add_paragraph('Ligação: viga parafusada à single plate; single plate soldada à '+('alma da viga de apoio.' if c.kind=='beam_web' else 'mesa do pilar. Não há solda direta entre viga e pilar.'))
    doc.add_paragraph('; '.join(mats)+'. E = 200 000 MPa; γa1 = 1,10; γa2 = γw2 = 1,35.')
    doc.add_paragraph('Esforços e hipóteses',style='Heading 1')
    doc.add_paragraph(f'Entrada já majorada: V = {number(c.V/KGF)} kgf ({number(c.V/1000,3)} kN); N = {number(c.N/KGF)} kgf ({number(c.N/1000,3)} kN), tração positiva. Não se aplicam novos coeficientes de ações.')
    if r.minimum_factor>1:
        doc.add_paragraph(f'Verificação adicional do mínimo de 45 kN conforme NBR 8800, 6.1.5.2, na direção da resultante informada: V = {number(c.V*r.minimum_factor/KGF)} kgf; N = {number(c.N*r.minimum_factor/KGF)} kgf. Os resultados governantes abaixo incluem essa verificação.')
    doc.add_paragraph(f'Encontro a 90°; {c.bolt_columns} coluna(s) de parafusos; soldagem em oficina. Face do apoio ao centro do grupo: e = {number(c.bolt_centroid_x)} mm. Excentricidade vertical de N: eN = {number(c.eccentric_n)} mm. Referência geométrica: M = |V|e + |N·eN|. Contenção longitudinal da viga apoiada: eficaz por hipótese fixa definida para esta versão. Não há cortante horizontal nem momento no eixo de menor inércia. N é força axial de tração, distinta de cortante horizontal.')
    cope_label={'none':'sem recorte','top':'mesa superior','both':'mesas superior e inferior'}[c.cope]
    doc.add_paragraph(f'Posição z = {number(c.plate_top)} mm: distância da face superior da mesa da viga apoiada ao topo do trecho parafusado, antes do recorte. Desnível entre vigas: {number(c.beam_level)} mm. Recorte: {cope_label}; dimensões superior/inferior/comprimento = {c.coped_top:g}/{c.coped_bottom:g}/{c.cope_length if c.cope!="none" else 0:g} mm. Folga de montagem = {c.clearance:g} mm; raio do envelope de montagem = {c.tool_radius:g} mm. Borda da solda reforçada: '+('sim.' if c.reinforced_weld else 'não.'))
    if c.kind=='column_flange':
        m,p,_=column_actions(c,c.V*r.minimum_factor,c.N*r.minimum_factor)
        doc.add_paragraph(f'Transmissão ao pilar: Mma = |V|(e+tf)+|N·eN| = {number(m/(KGF*1000))} kgf·m; resultantes locais T = {number(p.tension/KGF)} kgf e C = {number(p.compression/KGF)} kgf. Integração da carga linear p(y)=N/hp+12Mma·y/hp³; T−C=N. Inclui o mínimo normativo quando aplicável.')
        doc.add_paragraph('Hipótese do apoio: deslocamento lateral relativo entre as mesas do pilar impedido na região da ligação (NBR 8800, 5.7.5.1). A análise global do pilar permanece no projeto estrutural. Essa contenção é premissa, não conclusão decorrente da solda de penetração total.')
    if c.notes:doc.add_paragraph(c.notes)
    doc.add_paragraph('Referências e limites de aplicação',style='Heading 1')
    doc.add_paragraph('ABNT NBR 8800:2024, errata 2025: 5.2.4; 5.7; 6.1.5.2; 6.2; 6.3; 6.5 e Anexo A. AISC Companion v16.0, Volume 1, P901-23W: exemplos II.A-6, II.A-7, II.A-17B, II.A-18 e II.A-19B. SCI P358 (2014), seção 5. Muir e Hewitt (2009), Engineering Journal 46(2), p.67–80. Dowswell (2018), Engineering Journal 55(4), p.231–242. Procedimentos complementares AISC/SCI identificados por verificação, com resistências e coeficientes explicitados.')
    if c.full_depth:doc.add_paragraph('Chapas estendidas: Muir e Hewitt, Engineering Journal 46(2), 2009, p.67–80; Thornton e Fortney, Engineering Journal 48(2), 2011; Motallebi, Lignos e Rogers, Journal of Constructional Steel Research 148, 2018, p.336–350. Referências para delimitação do modelo e condição de compacidade; não constituem validação desta variante na versão atual.')
    if any(name.startswith('USI-CIVIL') for label,name in material_list):
        doc.add_paragraph('USI-CIVIL: propriedades mínimas e espessuras de 6 a 75 mm conforme Usiminas, Catálogo de chapas grossas, jul. 2022, p.29. Aplicação em chapas e perfis soldados; não são especificações de perfis W laminados.')
        hyperlink(doc.add_paragraph(),'Catálogo Usiminas — fonte das propriedades','https://usiminas.com/wp-content/uploads/2024/04/CatalagoChapasGrossas.pdf')
    doc.add_paragraph(f'Versão {VERSION}, escopo delimitado: uma viga, tração axial e cortante no plano. A chapa/enrijecedores entre mesas não são avaliados nesta versão. Não inclui fadiga, atrito, ações cíclicas ou incêndio. O mínimo normativo é avaliado pela resultante; a análise global das barras é externa.')
    if c.kind=='beam_web' and c.N>0 and not c.full_depth and not c.support_web_combined_excluded:
        doc.add_paragraph('Plastificação da alma: AISC Manual 16, eq.9-45; P901 II.A-19B, IIA-228; Kapp, Engineering Journal 11(2), 1974, p.38–41. A expressão é aplicada somente à tração direta centrada, sem cortante, e exige a distância longitudinal livre indicada.')
    if pending:
        doc.add_paragraph('Pendências para conclusão',style='Heading 2')
        for i in pending:doc.add_paragraph((i.origin_label+': ' if i.severity=='pending' else 'Geometria: ')+i.text)
    doc.add_paragraph(('Resumo dos componentes isolados' if c.full_depth else 'Resumo das verificações') if r.checks else 'Cobertura do pré-detalhamento',style='Heading 1')
    checkrows=[]
    for x in r.checks:
        sd,unit=display(x.demand,x.unit);rd,_=display(x.resistance,x.unit)
        checkrows.append((x.name,sd,rd,unit,number(x.ratio,3),'OK' if x.passed else 'NÃO'))
    if r.checks:
        table(doc,['Verificação','Sd / valor','Rd / limite','Unid.','Índice','Atende'],checkrows,[2.95,1,1,.65,.8,.7])
        doc.add_paragraph('Condições de método, incluindo compacidade, não são índices de utilização sob a carga aplicada. Índices de interação também não são multiplicadores lineares da carga. O atendimento usa valores sem arredondamento e abrange somente as verificações realizadas.')
        doc.add_paragraph('Memória de cálculo',style='Heading 1')
        doc.add_paragraph('As substituições usam N, mm e MPa; os resultados principais são apresentados em kgf e kgf·m. Conversão: 1 kgf = 9,80665 N. Variáveis geométricas seguem o desenho. Na versão compacta, desenvolvem-se as verificações determinantes por componente.')
    else:
        doc.add_paragraph('Conferidas as dimensões, o arranjo dos furos, os envelopes de montagem e os limites geométricos de solda implementados. Não foram verificadas resistências, redistribuição de esforços, estabilidade ou capacidade de rotação do conjunto. O mínimo normativo de força também não foi aplicado a resistências nesta variante.')
        doc.add_paragraph(f'Área geométrica da chapa, antes da dedução dos furos: {number(polygon_area(plate_outline(c)))} mm². Essa área serve apenas para detalhamento e não define uma área resistente equivalente.')
    if detailed:selected=r.checks
    elif c.full_depth:
        selected=[]
        for prefix in ('full_weld_','full_base_','full_compactness','opposite_compactness'):
            candidates=[x for x in r.checks if x.id.startswith(prefix)]
            if candidates:selected.append(max(candidates,key=lambda x:x.ratio))
    else:
        groups=[['bolts'],['bearing_plate','bearing_beam'],['plate_ltb','plate_mu','interaction_y','interaction_u','block_plate','block_plate_u','block_plate_partial_u','partial_plate_l'],['beam_v','beam_n','block_beam_u','beam_vm','beam_buckling','cope_interaction','cope_rupture','cope_block','block_beam_partial_u','partial_beam_l'],['weld','base_plate'],['support_shear','support_punch','support_web_y','support_flange','support_joint_weld','support_joint_base'],['support_crippling'],['ductility','weld_development']]
        selected=[]
        groups.extend([['support_web_n'],['support_web_punch_n']])
        for groupids in groups:
            candidates=[x for x in r.checks if x.id in groupids]
            if candidates:selected.append(max(candidates,key=lambda x:x.ratio))
    for x in selected:
        start=len(doc.paragraphs)
        doc.add_paragraph(x.name,style='Heading 2')
        native_equation(doc,x,c)
        if x.id=='cope_interaction':doc.add_paragraph(x.equation.split('; ',1)[1])
        doc.add_paragraph(decimal_text(x.substitution))
        sd,u=display(x.demand,x.unit);rd,_=display(x.resistance,x.unit)
        labels=('Valor','Limite') if x.category=='Condição do método' else ('Sd','Rd')
        doc.add_paragraph(f'Resultado: {labels[0]} = {sd} {u}; {labels[1]} = {rd} {u}; índice = {number(x.ratio,3)}. '+('Atende.' if x.passed else 'Não atende.'))
        doc.add_paragraph(x.variables+' Referência: '+x.reference)
        for paragraph in doc.paragraphs[start:-1]:
            paragraph.paragraph_format.keep_with_next=True
    p=s.footer.paragraphs[0];p.paragraph_format.space_after=Pt(0)
    r0=p.add_run(BRAND+' ');r0.font.size=Pt(7)
    hyperlink(p,'LinkedIn de Lucas Oliveira',LINK)
    p=s.footer.add_paragraph(f'LRO Ligações v{VERSION} · ');p.paragraph_format.space_after=Pt(0)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
    for r0 in p.runs:r0.font.size=Pt(7)
    buf=BytesIO();doc.save(buf);return buf.getvalue()
